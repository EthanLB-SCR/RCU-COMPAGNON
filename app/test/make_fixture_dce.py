#!/usr/bin/env python3
# make_fixture_dce.py — génère test/fixture_dce.zip (un DCE fictif : CCTP 2 pages avec articles numérotés, RC, DPGF Excel, 2 plans, un PPSPS Word,
# un zip imbriqué de profils) et test/fixture_cctp.pdf (le CCTP seul), pour les tests du dossier et de la lecture des pièces (e2e_commerce2).
import os, io, zipfile, shutil
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document
HERE = os.path.dirname(os.path.abspath(__file__))
def pdf(paragraphs, title):
    buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=A4); W, H = A4
    y = H - 60; c.setFont('Helvetica-Bold', 14); c.drawString(50, y, title); y -= 30
    for kind, text in paragraphs:
        if kind == 'page': c.showPage(); y = H - 60; continue
        if kind == 'h': c.setFont('Helvetica-Bold', 12); y -= 8
        else: c.setFont('Helvetica', 10.5)
        for line in wrap(text, 95 if kind != 'h' else 80):
            if y < 60: c.showPage(); y = H - 60; c.setFont('Helvetica-Bold' if kind == 'h' else 'Helvetica', 12 if kind == 'h' else 10.5)
            c.drawString(50, y, line); y -= 15
        y -= 6
    c.save(); return buf.getvalue()
def wrap(t, n):
    out, cur = [], ''
    for w in t.split():
        if len(cur) + len(w) + 1 > n: out.append(cur); cur = w
        else: cur = (cur + ' ' + w).strip()
    if cur: out.append(cur)
    return out
CCTP = [('h', '4.2 - Tubes et pieces pre-isolees'), ('p', "Les canalisations seront en acier pre-isole serie 2, conformes a l'EN 253, tubes de 12 m, gaine PEHD, fils d'alarme integres. Fournisseur agree par le maitre d'oeuvre. Les pieces speciales (coudes, tes, reductions, points fixes) sont pre-isolees d'usine."),
        ('h', '4.5 - Assemblage par soudage'), ('p', "Les soudures bout a bout seront executees par des soudeurs qualifies selon EN ISO 9606-1, procede 111 ou 141 selon DMOS valide. 10 % des soudures feront l'objet d'un controle radiographique, choix des joints par le maitre d'oeuvre. Toute soudure non conforme sera refaite aux frais de l'entreprise, controle a 100 % de la serie suivante."),
        ('h', '4.7 - Manchons et retraction'), ('p', "Les jonctions de gaine seront realisees par manchons electrosoudables avec controle d'etancheite a l'air (0,2 bar, 2 min) avant injection de mousse. Les fils d'alarme seront raccordes et testes a chaque jonction."),
        ('page', ''), ('h', '5.1 - Epreuves'), ('p', "Epreuve hydraulique a 1,5 fois la pression de service pendant 24 heures par troncon, PV contradictoire avec le maitre d'oeuvre. Le rincage et le remplissage en eau adoucie sont a la charge de l'entreprise."),
        ('h', '6.3 - Tranchee et remblais'), ('p', "La fourniture et la pose en tranchee ouverte incluent le lit de pose en sable 0/4, le remblai en sable jusqu'a 10 cm au-dessus de la generatrice superieure et le grillage avertisseur. Les terrassements sont au lot 1.")]
RC = [('h', '1 - Objet de la consultation'), ('p', "Extension du reseau de chaleur, lot 2 : reseau acier pre-isole. Procedure adaptee (MAPA), lots separes. Variante autorisee."), ('h', '2 - Dates'), ('p', "Visite de site obligatoire le 14/10 a 9 h. Questions jusqu'au 16/10. Remise des offres le 23/10 avant 12 h. Delai d'execution : 22 semaines."), ('h', '3 - Jugement des offres'), ('p', "Prix 60 %, valeur technique 40 %.")]
cctp = pdf(CCTP, 'CCTP - Lot 2 - Reseau acier pre-isole'); open(os.path.join(HERE, 'fixture_cctp.pdf'), 'wb').write(cctp)
rc = pdf(RC, 'Reglement de la consultation')
plan1 = pdf([('p', 'Plan PL-01 - trace general - echelle 1/500')], 'PL-01'); plan2 = pdf([('p', 'Plan PL-02 - profil en travers')], 'PL-02')
profil = pdf([('p', 'Profil en long - rue de Kerentrech')], 'Profil PL-01')
d = Document(); d.add_heading('PPSPS cadre', 1); d.add_heading('1. Organisation du chantier', 2); d.add_paragraph("Le chantier est base a la sous-station K1. Accueil securite de chaque intervenant par le chef de chantier. Port des EPI obligatoire.")
d.add_heading('2. Risques particuliers', 2); d.add_paragraph("Travaux en tranchee a proximite de reseaux enterres (DT-DICT). Soudage : risque incendie, permis de feu. Manutention de tubes de 12 m a la pelle.")
b = io.BytesIO(); d.save(b); ppsps = b.getvalue()
dpgf = open(os.path.join(HERE, 'fixture_dpgf.xlsx'), 'rb').read()
nested = io.BytesIO()
with zipfile.ZipFile(nested, 'w', zipfile.ZIP_DEFLATED) as z: z.writestr('Profils/Profil_PL-01.pdf', profil)
out = os.path.join(HERE, 'fixture_dce.zip')
with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('DCE lot 2/CCTP_lot2.pdf', cctp); z.writestr('DCE lot 2/RC.pdf', rc); z.writestr('DCE lot 2/DPGF_lot2.xlsx', dpgf)
    z.writestr('DCE lot 2/Plans/PL-01_trace.pdf', plan1); z.writestr('DCE lot 2/Plans/PL-02_profil.pdf', plan2); z.writestr('DCE lot 2/PPSPS_cadre.docx', ppsps)
    z.writestr('DCE lot 2/Profils.zip', nested.getvalue()); z.writestr('__MACOSX/._CCTP_lot2.pdf', b'junk'); z.writestr('DCE lot 2/.DS_Store', b'junk')
print('écrit', out, os.path.getsize(out), 'octets')
