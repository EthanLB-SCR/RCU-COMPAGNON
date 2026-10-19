# Fabrique test/fixture_dpgf_cols.xlsx — DPGF FICTIF « difficile » pour le lot DY (19/10 : lire n'importe quel DPGF client). Il reproduit les pièges d'un
# bordereau de client réel, sans aucune donnée réelle :
#   · des lignes de titre au-dessus de l'en-tête, dont une avec « Détail des prix » (l'ancienne lecture la prenait pour l'en-tête : tout devenait des titres) ;
#   · un en-tête sur 2 lignes (« Prix » sur « unitaire HT », « Quantité » sur « Marché ») et une colonne « U.M. » ;
#   · le PU en colonne O, des colonnes vides (C, D, F, H à N) ;
#   · des n° « 3.2.1 » ; des unités écrites de plusieurs façons (m.l., ML, mètre linéaire, m, U, Ens, Fft, m²) ;
#   · des quantités en formules (sans valeur en cache : le fichier est écrit par openpyxl) ; une ligne de prix sans unité (4.2).
# Zones fictives : « Zone Nord (essai) » (DN 100 : 2.1, 2.2, 2.3) et « Zone Sud (essai) » (DN 80 : 3.2.1, 3.2.2, 3.2.3).
# Lancer : python3 test/make_fixture_dpgfcols.py
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'fixture_dpgf_cols.xlsx')
EUR = '#,##0.00\\ "€"'
thin = Side(style='thin', color='808080'); med = Side(style='medium', color='000000')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); ws = wb.active; ws.title = 'DPGF lot 05'
widths = {'A': 9, 'B': 52, 'C': 2, 'D': 2, 'E': 8, 'F': 2, 'G': 11, 'O': 13, 'P': 15}
for col in 'HIJKLMN': widths[col] = 2
for col, w in widths.items(): ws.column_dimensions[col].width = w
for col in 'IJKLM': ws.column_dimensions[col].hidden = True
# titres
ws.merge_cells('A1:P1'); ws['A1'] = 'Réseau de chaleur des Essais — opération fictive (pièce de test)'
ws['A1'].font = Font(bold=True, size=14, color='FFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='1F3864'); ws['A1'].alignment = Alignment(horizontal='center')
ws.merge_cells('A2:P2'); ws['A2'] = 'DPGF · Détail des prix'; ws['A2'].font = Font(bold=True, size=12); ws['A2'].alignment = Alignment(horizontal='center')
ws['A3'] = 'Lot 05 — Réseau pré-isolé · aucune donnée réelle, quantités et désignations inventées'; ws['A3'].font = Font(italic=True, color='555555')
ws['A5'] = 'Montants en euros hors taxes'; ws['A5'].font = Font(italic=True, color='555555')
# en-tête sur 2 lignes
H1 = {'A': 'N°', 'B': 'Désignation des ouvrages', 'E': 'U.M.', 'G': 'Quantité', 'O': 'Prix', 'P': 'Prix'}
H2 = {'G': 'Marché', 'O': 'unitaire HT', 'P': 'total HT'}
for col in 'ABCDEFGHIJKLMNOP':
    for r, H in ((6, H1), (7, H2)):
        c = ws[f'{col}{r}']
        if col in H: c.value = H[col]
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9D9D9'); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(left=thin, right=thin, top=med if r == 6 else None, bottom=med if r == 7 else None)
ws.merge_cells('A6:A7'); ws.merge_cells('B6:D7'); ws.merge_cells('E6:E7')
rows = [
    ('1', 'INSTALLATIONS DE CHANTIER', None, None, 'chap'),
    ('1.1', 'Installation de chantier, base vie et repli', 'Fft', 1, 'line'),
    ('2', 'RÉSEAU — ZONE NORD (essai)', None, None, 'chap'),
    ('2.1', 'Tube acier pré-isolé DN 100 — rue des Essais', 'm.l.', '=120+60', 'line'),
    ('2.2', 'Tube acier pré-isolé DN 100 — traversée du parc fictif', 'ML', '=2*60', 'line'),
    ('2.3', 'Tube acier pré-isolé DN 100 — extension (essai)', 'ml', 48, 'line'),
    ('2.4', 'Soudure bout à bout DN 100 en fouille', 'U', 6, 'line'),
    ('3', 'RÉSEAU — ZONE SUD (essai)', None, None, 'chap'),
    ('3.1', 'Réfection provisoire de chaussée', 'm²', 35, 'line'),
    ('3.2', 'Canalisations DN 80', None, None, 'chap'),
    ('3.2.1', 'Tube acier pré-isolé DN 80 — allée des Tests', 'ml', '=100+60', 'line'),
    ('3.2.2', 'Tube acier pré-isolé DN 80 — antenne bâtiment B', 'mètre linéaire', 80, 'line'),
    ('3.2.3', 'Tube acier pré-isolé DN 80 — antenne bâtiment C (essai)', 'm', 20, 'line'),
    ('4', 'ESSAIS ET DIVERS', None, None, 'chap'),
    ('4.1', 'Épreuve hydraulique du réseau', 'Ens', 1, 'line'),
    ('4.2', 'Divers et sujétions d\'exécution', None, 1, 'line'),
]
r = 8; first = r
for n, des, u, q, kind in rows:
    ws[f'A{r}'] = n; ws[f'B{r}'] = des
    for col in 'ABCDEFGHIJKLMNOP': ws[f'{col}{r}'].border = box
    if kind == 'chap':
        for col in 'ABCDEFGHIJKLMNOP': ws[f'{col}{r}'].fill = PatternFill('solid', fgColor='F2F2F2'); ws[f'{col}{r}'].font = Font(bold=True)
    else:
        if u is not None: ws[f'E{r}'] = u
        ws[f'E{r}'].alignment = Alignment(horizontal='center')
        ws[f'G{r}'] = q; ws[f'G{r}'].number_format = '#,##0.00'
        ws[f'O{r}'].number_format = EUR
        ws[f'P{r}'] = f'=G{r}*O{r}'; ws[f'P{r}'].number_format = EUR
        ws[f'B{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
last = r - 1
ws[f'B{r}'] = 'Sous-total du lot'; ws[f'B{r}'].font = Font(bold=True, italic=True)
ws[f'P{r}'] = f'=SUBTOTAL(9,P{first}:P{last})'; ws[f'P{r}'].number_format = EUR
r += 2
ws[f'B{r}'] = 'TOTAL HT'; ws[f'B{r}'].font = Font(bold=True); ws[f'P{r}'] = f'=P{r-2}'; ws[f'P{r}'].number_format = EUR
ws[f'B{r+1}'] = 'TVA 20 %'; ws[f'P{r+1}'] = f'=ROUND(P{r}*0.2,2)'; ws[f'P{r+1}'].number_format = EUR
ws[f'B{r+2}'] = 'TOTAL TTC'; ws[f'B{r+2}'].font = Font(bold=True); ws[f'P{r+2}'] = f'=P{r}+P{r+1}'; ws[f'P{r+2}'].number_format = EUR
ws.freeze_panes = 'A8'
wb.save(OUT)
print('écrit', OUT, '· lignes', first, 'à', last)
