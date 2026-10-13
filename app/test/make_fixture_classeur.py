#!/usr/bin/env python3
# make_fixture_classeur.py — génère test/fixture_classeur.xlsx : la MAQUETTE SYNTHÉTIQUE du classeur (fixture_prix.xlsx, prix fictifs) + un onglet d'étude
# « ZONE A » REMPLI (mêmes cellules que l'onglet « Soudure » du classeur V42) + un « Récap Prix » (une colonne d'étude). Étape 2 : node test/make_fixture_classeur.mjs
# écrit dans le Récap les montants « avec marge » que donnerait le classeur (moteur de l'appli + règles du classeur : jours CALENDAIRES pour la base vie,
# soudures sans case ignorées au contrôle) et un écart volontaire (+120 € sur les purges) — pour tester le contrôle croisé (lot 5, 13/10/2026).
# Le vrai classeur et les vrais prix ne sont jamais dans le dépôt.
# python3 test/make_fixture_classeur.py && node test/make_fixture_classeur.mjs
import os, datetime
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(HERE, 'fixture_prix.xlsx'))
DNS = [20, 25, 32, 40, 50, 65, 80, 100, 125, 150, 200, 250, 300, 350, 400, 450, 500, 600]
z = wb.create_sheet('ZONE A')
z['B2'] = 'ZONE A'
z['B19'] = 'SOUDURES - MANCHONS - AUTOCONTROLE'
# chantier
z['B9'] = 'Date travaux'; z['C9'] = datetime.date(2026, 11, 2); z['D9'] = datetime.date(2026, 11, 27); z['E9'] = 'Type chantier'; z['F9'] = 'Gros'
z['I9'] = 'Isolation'; z['I10'] = 'Iso 1'; z['J9'] = 'Nbre'; z['J10'] = 1
for a, v in {'B11': 'Oui', 'C11': 'Oui', 'D11': 'Non', 'B13': 'Oui', 'C13': 'Non', 'D13': 0, 'B15': 'Non', 'C15': 'Oui', 'D15': 1,
             'H13': 'Non', 'I13': 'S+T+P', 'J13': 'M', 'H15': 'Non', 'I15': 'Non', 'I16': 'Non', 'J16': 5}.items():
    z[a] = v
# équipes, manchons, autocontrôle
for a, v in {'D23': 'S+T+P', 'D24': 'M', 'D25': 'Thermo', 'G24': 'Oui', 'H24': 'Non', 'I24': 'Non', 'J24': 'Non', 'K24': 'Non', 'L24': 'Non', 'N24': 'Non', 'R24': False, 'V47': 'Non'}.items():
    z[a] = v
# table des DN : DN 100 feeder, DN 50 antenne, DN 32 sans case
z['B27'] = 'DN'; z['H27'] = 'Antenne'; z['I27'] = 'Feeder'; z['J27'] = 'Nombre de soudures'; z['L27'] = 'Ml de tranchée'
for i, dn in enumerate(DNS):
    r = 28 + i; z.cell(r, 2, dn); z.cell(r, 8, False); z.cell(r, 9, False)
z['J35'] = 40; z['L35'] = 200; z['I35'] = True
z['J32'] = 12; z['L32'] = 60; z['H32'] = True
z['J30'] = 6; z['L30'] = 10
# nuit (off), purges, vidanges
for a, v in {'C54': 100, 'C58': 0.8, 'H54': 32, 'I54': 'Oui', 'J54': 2, 'H57': 'Non', 'N54': 50, 'O54': 'VIDANGES PAR GRAVITE (VANNES PREISOLEES)', 'P54': 1, 'N58': 'Non'}.items():
    z[a] = v
# vannes préisolées, SST, extrusion, raccordement
for a, v in {'B67': 100, 'C67': 2, 'D67': 'Non', 'E67': 'Non', 'F67': 400, 'J67': 80, 'K67': 'VANNES BS PN25', 'M67': 2, 'S67': 110, 'T67': 'Manchons', 'U67': 2,
             'L85': 100, 'M85': 'Insertion de tés', 'Q85': 1}.items():
    z[a] = v
# CND : distance (calculée par le classeur depuis latitude / longitude), radio 10 % antenne et feeder
for a, v in {'I84': 120, 'E88': 20, 'F88': 0.1, 'E90': 20, 'F90': 0.1}.items():
    z[a] = v
# hydraulique : zone 1, épreuve seule
for a, v in {'B117': 100, 'C117': 400, 'C130': 1, 'E130': 4, 'G130': 'Non', 'C132': 'OS', 'C138': 'Non', 'C146': 'Non', 'C154': 'Non', 'C162': 'Non', 'M135': 'Non'}.items():
    z[a] = v
# aérien, dépose, panier sur mesure, fourniture préisolée (prix saisi)
for a, v in {'N114': 'Non', 'L119': 'Tradi', 'M119': 80, 'N119': 12, 'O119': 2, 'P119': 4, 'Q119': 'PVC', 'S119': 'Mural', 'R128': 0,
             'L149': 100, 'M149': 30, 'B175': 1, 'C175': 50, 'D175': 'Coudes à souder', 'E175': 3, 'F175': 7, 'I175': 1, 'J175': 'S+T', 'K175': 4, 'O10': 15000, 'P10': 0}.items():
    z[a] = v
# Récap Prix : une colonne d'étude (montants « avec marge », remplis à l'étape 2)
R = wb.create_sheet('Récap Prix')
for c, h in enumerate(['Désignation', '', '', 'Hors €/ml', 'ZONE A', 'Prix achat', 'RCFA', 'Marge', 'Prix Vente', 'Observations'], 1):
    if h: R.cell(1, c, h)
R['E2'] = 'Montant avec marge'; R['G3'] = 'Aucun'; R['H3'] = 0.125
LAB = ['Chantier et linéaire', 'Amené - Repli - Nettoyage chantier', 'Contener-Benne-Pelle', 'Déchargement', 'Transfert', 'Déduction bardage', 'Prestation',
       'Soudure (compris manutention)', 'Manchons', 'Purges', 'Vidanges', 'Vannes préisolées', 'Vannes SST', 'Raccordement existant', 'Contrôles - Essais',
       'Autocontrole interne', 'Coussins', 'Ressuage externe', 'Contrôles ultrasons', 'Radio', 'Epreuve hydraulique', 'Eau Adoucie', 'Boulet mousse', 'Rinçage',
       'Passivation', 'Mise en conservation', 'Autres', '- Fournitures pré isolées', '- Fournitures acier nu aérien', '- Fournitures autres : ',
       '- Calo + peinture + supp aérien', '- Pose aérien', '- Fonçage', '- Dépose réseau', '- Extrusion', '- Etiquettes', '- Plus value nuit',
       '- DOE / Plan iso si aérien / Carnet soudure si prévu', '- Prestations autres : ', '- Autres frais divers', 'Total']
for i, lab in enumerate(LAB):
    r = 4 + i; R.cell(r, 1, lab)
    if lab in ('Prestation', 'Contrôles - Essais', 'Autres', 'Chantier et linéaire', 'Total'): continue
    R.cell(r, 5, 0); R.cell(r, 6, 0); R.cell(r, 7, 0); R.cell(r, 8, 0.125); R.cell(r, 9, 0)
R['E4'] = 0; R['E43'] = 200  # autres frais divers saisis dans le Récap (avec marge)
R['E44'] = 0; R['F44'] = 0; R['G44'] = 0; R['H44'] = 0.125; R['I44'] = 0
OUT = os.path.join(HERE, 'fixture_classeur.xlsx')
wb.save(OUT); print('écrit', OUT, '(étape 2 : node test/make_fixture_classeur.mjs)')
