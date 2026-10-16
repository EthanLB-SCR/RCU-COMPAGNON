# Fabrique test/fixture_dpgf_clair.xlsx — DPGF FICTIF pour le lot C2 (16/10, DPGF plus clair) : forfaits d'installation, lignes au ml (DN 100 × 2, DN 80,
# 2 antennes DN 50), lignes à l'unité (soudures DN 100, vannes, purges), et des lignes d'AUTOCONTRÔLE comme dans un vrai DPGF (« Contrôle Itopo des manchons
# DN 100 », « Essai de dépression des manchons DN 100 », « Autocontrôles des manchons DN 80 »), épreuve, divers. Montants = Qté × PU, sous-totaux SOUS.TOTAL,
# total HT, TVA, TTC. Aucune donnée réelle. Lancer : python3 test/make_fixture_dpgfclair.py
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'fixture_dpgf_clair.xlsx')
EUR = '#,##0.00\\ "€"'
thin = Side(style='thin', color='808080'); med = Side(style='medium', color='000000')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); ws = wb.active; ws.title = 'DPGF'
for col, w in zip('ABCDEF', [8, 60, 7, 11, 13, 16]): ws.column_dimensions[col].width = w
ws.merge_cells('A1:F1'); ws['A1'] = 'DÉCOMPOSITION DU PRIX GLOBAL ET FORFAITAIRE'
ws['A1'].font = Font(bold=True, size=14, color='FFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='1F3864'); ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 28
ws.merge_cells('A2:F2'); ws['A2'] = 'Lot 04 — Réseau de chaleur · opération fictive (essai du DPGF plus clair)'; ws['A2'].font = Font(italic=True, color='555555'); ws['A2'].alignment = Alignment(horizontal='center')
ws.row_dimensions[3].height = 6
for i, h in enumerate(['N°', 'Désignation', 'U', 'Qté', 'PU HT', 'Montant HT']):
    c = ws.cell(row=4, column=i + 1, value=h); c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9D9D9')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(left=thin, right=thin, top=med, bottom=med)
ws.row_dimensions[4].height = 30
rows = [
    ('1', 'INSTALLATION DE CHANTIER', None, None, 'chap'),
    ('1.1', 'Installation de chantier, base vie et repli', 'ft', 1, 'line'),
    ('1.2', 'Amené et repli du matériel de soudage', 'ft', 1, 'line'),
    ('', 'Sous-total chapitre 1', None, None, 'sub'),
    ('2', 'RÉSEAU PRÉ-ISOLÉ', None, None, 'chap'),
    ('2.1', 'Fourniture et pose tube acier pré-isolé DN 100 — rue des Essais', 'ml', 420, 'line'),
    ('2.2', 'Fourniture et pose tube acier pré-isolé DN 100 — traversée du parc', 'ml', 180, 'line'),
    ('2.3', 'Fourniture et pose tube acier pré-isolé DN 80 — avenue Fictive', 'ml', 300, 'line'),
    ('2.4', 'Tube pré-isolé DN50/125 — antenne bâtiment A', 'ml', 120, 'line'),
    ('2.5', 'Tube pré-isolé DN50/125 — antenne bâtiment B', 'ml', 80, 'line'),
    ('', 'Sous-total chapitre 2', None, None, 'sub'),
    ('3', 'OUVRAGES À L\'UNITÉ', None, None, 'chap'),
    ('3.1', 'Soudure bout à bout DN 100 en fouille', 'u', 40, 'line'),
    ('3.2', 'Vanne d\'isolement préisolée DN 80', 'u', 2, 'line'),
    ('3.3', 'Purges et vidanges, y compris robinetterie', 'ens.', 1, 'line'),
    ('', 'Sous-total chapitre 3', None, None, 'sub'),
    ('4', 'ESSAIS ET AUTOCONTRÔLES', None, None, 'chap'),
    ('4.1', 'Épreuve hydraulique du réseau', 'ft', 1, 'line'),
    ('4.2', 'Contrôle Itopo des manchons DN 100', 'u', 120, 'line'),
    ('4.3', 'Essai de dépression des manchons DN 100', 'u', 120, 'line'),
    ('4.4', 'Autocontrôles des manchons DN 80 (tous types)', 'u', 50, 'line'),
    ('', 'Sous-total chapitre 4', None, None, 'sub'),
    ('5', 'DIVERS', None, None, 'chap'),
    ('5.1', 'Divers et sujétions d\'exécution', 'ft', 1, 'line'),
    ('', 'Sous-total chapitre 5', None, None, 'sub'),
]
r = 5; subs = []; first = None
for n, des, u, q, kind in rows:
    ws.cell(row=r, column=1, value=n if n else None).alignment = Alignment(horizontal='left')
    ws.cell(row=r, column=2, value=des)
    for c in range(1, 7): ws.cell(row=r, column=c).border = box
    if kind == 'chap':
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        for c in range(1, 7): ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor='F2F2F2'); ws.cell(row=r, column=c).font = Font(bold=True)
        first = r + 1
    elif kind == 'line':
        ws.cell(row=r, column=3, value=u).alignment = Alignment(horizontal='center')
        ws.cell(row=r, column=4, value=q).number_format = '#,##0.00'
        ws.cell(row=r, column=5).number_format = EUR
        ws.cell(row=r, column=6, value=f'=D{r}*E{r}').number_format = EUR
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    else:
        ws.cell(row=r, column=2).font = Font(bold=True, italic=True)
        ws.cell(row=r, column=6, value=f'=SUBTOTAL(9,F{first}:F{r-1})').number_format = EUR; ws.cell(row=r, column=6).font = Font(bold=True)
        subs.append(r)
    r += 1
r += 1
tot = r
ws.cell(row=tot, column=2, value='TOTAL HT').font = Font(bold=True)
ws.cell(row=tot, column=6, value='=' + '+'.join(f'F{s}' for s in subs)).number_format = EUR
ws.cell(row=tot, column=6).font = Font(bold=True)
ws.cell(row=tot + 1, column=2, value='TVA')
ws.cell(row=tot + 1, column=5, value=0.2).number_format = '0%'
ws.cell(row=tot + 1, column=6, value=f'=ROUND(F{tot}*E{tot+1},2)').number_format = EUR
ws.cell(row=tot + 2, column=2, value='TOTAL TTC').font = Font(bold=True, color='1F3864')
ws.cell(row=tot + 2, column=6, value=f'=F{tot}+F{tot+1}').number_format = EUR
ws.freeze_panes = 'A5'
wb.save(OUT)
print('écrit', OUT, '· total HT ligne', tot)
