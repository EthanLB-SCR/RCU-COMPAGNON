# Fabrique test/fixture_dpgf_xl.xlsx — DPGF FICTIF « comme un vrai » pour la grille Excel du DPGF (14/10 après-midi, lot C).
# Deux feuilles (DPGF, Récap), fusions, styles (gras, italique, fonds, bordures), formats (€, %, décimales, comptabilité), volets figés,
# colonne masquée, formules : produit, SOUS.TOTAL, ARRONDI, SI, MAX, renvois entre feuilles, nom défini, une formule PARTAGÉE (comme Excel),
# une fonction non gérée par l'appli (ECARTYPE → « non recalculée »). Valeurs en cache écrites comme Excel (0 : PU vides).
# Aucune donnée réelle. Lancer : python3 test/make_fixture_dpgfxl.py
import os, re, zipfile, shutil, tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'fixture_dpgf_xl.xlsx')
EUR = '#,##0.00\\ "€"'
ACC = '_-* #,##0.00\\ "€"_-;\\-* #,##0.00\\ "€"_-;_-* "-"??\\ "€"_-;_-@_-'
thin = Side(style='thin', color='808080'); med = Side(style='medium', color='000000')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); ws = wb.active; ws.title = 'DPGF'
for col, w in zip('ABCDEFG', [8, 52, 7, 11, 13, 16, 14]): ws.column_dimensions[col].width = w
ws.column_dimensions['G'].hidden = True
ws.merge_cells('A1:F1'); ws['A1'] = 'DÉCOMPOSITION DU PRIX GLOBAL ET FORFAITAIRE'
ws['A1'].font = Font(bold=True, size=14, color='FFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='1F3864'); ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 28
ws.merge_cells('A2:F2'); ws['A2'] = 'Lot 02 — Réseau de chaleur · opération fictive (essai)'; ws['A2'].font = Font(italic=True, color='555555'); ws['A2'].alignment = Alignment(horizontal='center')
ws.row_dimensions[3].height = 6
hdr = ['N°', 'Désignation', 'U', 'Qté', 'PU HT', 'Montant HT', 'Code interne']
for i, h in enumerate(hdr):
    c = ws.cell(row=4, column=i + 1, value=h); c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9D9D9')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(left=thin, right=thin, top=med, bottom=med)
ws.row_dimensions[4].height = 30
rows = [
    ('1', 'INSTALLATION DE CHANTIER', None, None, 'chap'),
    ('1.1', 'Installation de chantier, base vie et repli', 'ft', 1, 'line'),
    ('1.2', "Plans d'exécution et dossier des ouvrages exécutés", 'ft', 1, 'line'),
    ('', 'Sous-total chapitre 1', None, None, 'sub'),
    ('2', 'RÉSEAU PRÉ-ISOLÉ', None, None, 'chap'),
    ('2.1', 'Fourniture et pose tube acier pré-isolé série 2 DN 100', 'ml', 600, 'line'),
    ('2.2', 'Soudure bout à bout DN 100, contrôle visuel', 'u', 100, 'line'),
    ('2.3', 'Manchon de jonction thermorétractable DN 100', 'u', 100, 'line'),
    ('2.4', 'Purges et vidanges, y compris vannes', 'ens.', 1, 'line'),
    ('', 'Sous-total chapitre 2', None, None, 'sub'),
    ('3', 'ESSAIS ET CONTRÔLES', None, None, 'chap'),
    ('3.1', 'Épreuve hydraulique du réseau', 'ft', 1, 'line'),
    ('3.2', 'Contrôle radiographique des soudures (10 %)', 'u', 10, 'line'),
    ('', 'Sous-total chapitre 3', None, None, 'sub'),
]
r = 5; subs = []; first = None
for n, des, u, q, kind in rows:
    ws.cell(row=r, column=1, value=n if n else None).alignment = Alignment(horizontal='left')
    ws.cell(row=r, column=2, value=des)
    for c in range(1, 7): ws.cell(row=r, column=c).border = box
    if kind == 'chap':
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        for c in range(1, 7): ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor='F2F2F2'); ws.cell(row=r, column=c).font = Font(bold=True)
        first = r + 1
    elif kind == 'line':
        ws.cell(row=r, column=3, value=u).alignment = Alignment(horizontal='center')
        ws.cell(row=r, column=4, value=q).number_format = '#,##0.00'
        ws.cell(row=r, column=5).number_format = EUR
        ws.cell(row=r, column=6, value=f'=ROUND(D{r}*E{r},2)' if n == '3.1' else f'=D{r}*E{r}').number_format = ACC
        ws.cell(row=r, column=7, value='INT-' + n.replace('.', '-'))
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    else:
        ws.cell(row=r, column=2).font = Font(bold=True, italic=True)
        ws.cell(row=r, column=6, value=f'=SUBTOTAL(9,F{first}:F{r-1})').number_format = ACC; ws.cell(row=r, column=6).font = Font(bold=True)
        subs.append(r)
    r += 1
r += 1
tot = r
ws.cell(row=tot, column=2, value='TOTAL HT').font = Font(bold=True)
ws.cell(row=tot, column=6, value='=' + '+'.join(f'F{s}' for s in subs)).number_format = ACC
ws.cell(row=tot, column=6).font = Font(bold=True); ws.cell(row=tot, column=6).border = Border(top=med, bottom=thin, left=thin, right=thin)
ws.cell(row=tot + 1, column=2, value='TVA')
ws.cell(row=tot + 1, column=5, value=0.2).number_format = '0%'
ws.cell(row=tot + 1, column=6, value=f'=ROUND(F{tot}*E{tot+1},2)').number_format = ACC
ws.cell(row=tot + 2, column=2, value='TOTAL TTC').font = Font(bold=True, color='1F3864')
ws.cell(row=tot + 2, column=6, value=f'=F{tot}+F{tot+1}').number_format = ACC
for c in range(2, 7): ws.cell(row=tot + 2, column=c).fill = PatternFill('solid', fgColor='DDEBF7')
ws.freeze_panes = 'A5'
rc = wb.create_sheet('Récap')
rc.column_dimensions['A'].width = 26; rc.column_dimensions['B'].width = 18
rc.merge_cells('A1:B1'); rc['A1'] = 'RÉCAPITULATIF'; rc['A1'].font = Font(bold=True, size=12); rc['A1'].alignment = Alignment(horizontal='center')
rc['A3'] = 'Chapitre'; rc['B3'] = 'Montant HT'
for c in ('A3', 'B3'): rc[c].font = Font(bold=True); rc[c].fill = PatternFill('solid', fgColor='D9D9D9')
rc['A4'] = '1 - Installation'; rc['B4'] = f'=DPGF!F{subs[0]}'
rc['A5'] = '2 - Réseau'; rc['B5'] = f"='DPGF'!F{subs[1]}"
rc['A6'] = '3 - Essais'; rc['B6'] = f'=DPGF!F{subs[2]}'
rc['A7'] = 'Total HT'; rc['B7'] = '=SUM(B4:B6)'; rc['A7'].font = Font(bold=True); rc['B7'].font = Font(bold=True)
rc['A8'] = 'Part du réseau'; rc['B8'] = '=IF(B7>0,B5/B7,0)'; rc['B8'].number_format = '0.0%'
rc['A9'] = 'Plus gros chapitre'; rc['B9'] = '=MAX(B4:B6)'
rc['A10'] = 'Contrôle'; rc['B10'] = f'=IF(ABS(B7-DPGF!F{tot})<0.01,"OK","ÉCART")'
rc['A11'] = 'Total HT (nom défini)'; rc['B11'] = '=TotalHT*1'
rc['A12'] = 'Écart type (non géré)'; rc['B12'] = '=STDEV(B4:B6)'
for c in ('B4', 'B5', 'B6', 'B7', 'B9', 'B11', 'B12'): rc[c].number_format = EUR
wb.defined_names['TotalHT'] = DefinedName('TotalHT', attr_text="'Récap'!$B$7")
tmp = tempfile.mkdtemp(); raw = os.path.join(tmp, 'raw.xlsx'); wb.save(raw)
# post-traitement : formule partagée F{lignes du chapitre 2} (comme l'écrit Excel) + valeurs en cache (0, « OK »)
z = zipfile.ZipFile(raw); files = {n: z.read(n) for n in z.namelist()}; z.close()
s1 = files['xl/worksheets/sheet1.xml'].decode('utf-8')
ch2 = [5 + i for i, x in enumerate(rows) if x[4] == 'line' and x[0].startswith('2.')]
a, b = ch2[0], ch2[-1]
def setf(xml, ref, inner):
    return re.sub(r'(<c r="%s"[^>]*>)<f>[^<]*</f>(<v\s*/>|<v>[^<]*</v>)?' % ref, lambda m: m.group(1) + inner, xml)
s1 = setf(s1, f'F{a}', f'<f t="shared" ref="F{a}:F{b}" si="0">D{a}*E{a}</f><v>0</v>')
for rr in range(a + 1, b + 1): s1 = setf(s1, f'F{rr}', '<f t="shared" si="0"/><v>0</v>')
s1 = re.sub(r'(<f>[^<]*</f>)(?:<v\s*/>|<v></v>)?(?=</c>)', r'\1<v>0</v>', s1)
files['xl/worksheets/sheet1.xml'] = s1.encode('utf-8')
s2 = files['xl/worksheets/sheet2.xml'].decode('utf-8')
s2 = re.sub(r'(<c r="B10"[^>]*?)(\s*t="\w+")?>(<f>[^<]*</f>)(?:<v\s*/>|<v></v>)?', r'\1 t="str">\3<v>OK</v>', s2)
s2 = re.sub(r'(<f>[^<]*</f>)(?:<v\s*/>|<v></v>)?(?=</c>)', r'\1<v>0</v>', s2)
files['xl/worksheets/sheet2.xml'] = s2.encode('utf-8')
# textes partagés (comme Excel) au lieu des textes en ligne d'openpyxl, et une chaîne de calcul (xl/calcChain.xml) comme en laisse Excel
import html
sst = []; idx = {}
def to_shared(xml):
    def rep(m):
        t = html.unescape(m.group(2))
        if t not in idx: idx[t] = len(sst); sst.append(t)
        return m.group(1) + ' t="s"><v>%d</v></c>' % idx[t]
    return re.sub(r'(<c r="[A-Z]+\d+"(?: s="\d+")?) t="inlineStr"><is><t>([^<]*)</t></is></c>', rep, xml)
chain = []
for i, n in ((1, 'xl/worksheets/sheet1.xml'), (2, 'xl/worksheets/sheet2.xml')):
    x = to_shared(files[n].decode('utf-8')); files[n] = x.encode('utf-8')
    for k, ref in enumerate(re.findall(r'<c r="([A-Z]+\d+)"[^>]*><f', x)): chain.append('<c r="%s"%s/>' % (ref, ' i="%d"' % i if k == 0 else ''))
esc = lambda t: t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
files['xl/sharedStrings.xml'] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d" uniqueCount="%d">' % (len(sst), len(sst)) + ''.join('<si><t xml:space="preserve">%s</t></si>' % esc(t) for t in sst) + '</sst>').encode('utf-8')
files['xl/calcChain.xml'] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">' + ''.join(chain) + '</calcChain>').encode('utf-8')
ct = files['[Content_Types].xml'].decode('utf-8').replace('</Types>', '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/><Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
files['[Content_Types].xml'] = ct.encode('utf-8')
rl = files['xl/_rels/workbook.xml.rels'].decode('utf-8').replace('</Relationships>', '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rId5"/><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml" Id="rId6"/></Relationships>')
files['xl/_rels/workbook.xml.rels'] = rl.encode('utf-8')
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zo:
    for n, d in files.items(): zo.writestr(n, d)
shutil.rmtree(tmp)
print('écrit', OUT, '· lignes du chapitre 2 :', a, '→', b, '· total HT ligne', tot)
