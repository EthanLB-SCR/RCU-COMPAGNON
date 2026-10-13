# Fixtures « Word » de la Lecture (13/10/2026) : un CCTP fictif en .docx (titres NUMÉROTÉS par le style, comme les vrais CCTP : 1. / 2.1…, liste à puces,
# tableau, image, tabulation, paragraphe de titre vide qui ne doit pas consommer de numéro), puis le MÊME document converti par LibreOffice en .doc (Word 97-2003),
# .rtf et .odt. Lancer depuis app/test : python3 make_fixture_word.py (LibreOffice requis). Aucune donnée réelle.
import os, io, subprocess, shutil, struct, zlib
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
here = os.path.dirname(os.path.abspath(__file__))
def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')
doc = Document()
# numérotation multiniveau « 1. » / « 1.1 » attachée aux styles Titre 1 / Titre 2 (comme dans les CCTP)
numbering = doc.part.numbering_part.element
absn = OxmlElement('w:abstractNum'); absn.set(qn('w:abstractNumId'), '90')
for lvl, txt in ((0, '%1.'), (1, '%1.%2')):
    l = OxmlElement('w:lvl'); l.set(qn('w:ilvl'), str(lvl))
    for tag, val in (('w:start', '1'), ('w:numFmt', 'decimal'), ('w:lvlText', txt)):
        e = OxmlElement(tag); e.set(qn('w:val'), val); l.append(e)
    absn.append(l)
numbering.insert(0, absn)
num = OxmlElement('w:num'); num.set(qn('w:numId'), '90'); an = OxmlElement('w:abstractNumId'); an.set(qn('w:val'), '90'); num.append(an); numbering.append(num)
for name, lvl in (('Heading 1', 0), ('Heading 2', 1)):
    st = doc.styles[name]; pPr = st.element.get_or_add_pPr(); numPr = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(lvl)); ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '90'); numPr.append(il); numPr.append(ni); pPr.insert(0, numPr)
doc.add_heading('CCTP fictif — Lot 2 réseau de chaleur', 0)
doc.add_heading('Généralités', 1)
doc.add_paragraph('Le présent CCTP décrit les travaux de soudure et de manchonnage du lot 2 (fixture de test, aucune donnée réelle).')
doc.add_heading('', 1)  # titre vide : ne doit pas décaler la numérotation
doc.add_heading('Contrôles', 1)
doc.add_heading('Radiographie', 2)
doc.add_paragraph('Les soudures seront contrôlées par radiographie à 10 % des soudures sur les feeders et 20 % sur les antennes.')
p = doc.add_paragraph(); r = p.add_run('Vacation :'); r.bold = True; p.add_run('\t9 soudures par vacation, forfait déplacement compris.')
doc.add_heading('Épreuve hydraulique', 2)
doc.add_paragraph('Épreuve à 1,5 fois la pression de service, tronçon par tronçon.')
for t in ('Remplissage en eau adoucie', 'Rinçage dynamique', 'Mise en conservation'): doc.add_paragraph(t, style='List Bullet')
tb = doc.add_table(rows=2, cols=3)
for i, v in enumerate(('DN', 'ml', 'Soudures', '100', '250', '42')): tb.cell(i // 3, i % 3).text = v
doc.add_picture(io.BytesIO(png(40, 20, (200, 60, 30))))
doc.add_heading('Réception', 1)
doc.add_paragraph('Les PV d\'épreuve sont remis au maître d\'œuvre avant la mise en service.')
out = os.path.join(here, 'fixture_word.docx'); doc.save(out)
tmp = os.path.join(here, '_lo_tmp'); os.makedirs(tmp, exist_ok=True)
for ext in ('doc', 'rtf', 'odt'):
    subprocess.run(['soffice', '-env:UserInstallation=file:///tmp/lo_fixture_word', '--headless', '--convert-to', ext, '--outdir', tmp, out], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    shutil.move(os.path.join(tmp, 'fixture_word.' + ext), os.path.join(here, 'fixture_word.' + ext))
shutil.rmtree(tmp, ignore_errors=True)
print('ok', [f for f in os.listdir(here) if f.startswith('fixture_word')])
